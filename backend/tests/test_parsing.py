import io
import zipfile

import pytest

from app.services.resume_parser import ParseStatus, detect_format, parse_resume, sanitize_filename
from tests.conftest import make_docx, make_image_only_pdf, make_pdf

LIMIT = 1024 * 1024


def test_valid_pdf():
    result = parse_resume("resume.pdf", make_pdf(), max_bytes=LIMIT)
    assert result.status == ParseStatus.PARSED
    assert result.detected_format == "pdf"
    assert "Jordan Example" in result.text
    assert result.page_count == 1 and result.page_offsets == [0]
    assert len(result.file_hash) == 64


def test_valid_docx():
    result = parse_resume("resume.docx", make_docx(), max_bytes=LIMIT)
    assert result.status == ParseStatus.PARSED
    assert result.detected_format == "docx"
    assert "Built FastAPI services" in result.text


def test_docx_tables_are_extracted():
    from docx import Document

    document = Document()
    document.add_paragraph("Alex Tablerow")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "Five years building FastAPI services"
    buf = io.BytesIO()
    document.save(buf)
    result = parse_resume("t.docx", buf.getvalue(), max_bytes=LIMIT)
    assert result.ok and "Python | Five years building FastAPI services" in result.text


def test_corrupted_pdf():
    data = b"%PDF-1.7\n" + bytes(range(256)) * 3
    result = parse_resume("broken.pdf", data, max_bytes=LIMIT)
    assert result.status == ParseStatus.FAILED
    assert result.error_code == "corrupted_file"


def test_corrupted_docx():
    data = make_docx()[:400]
    result = parse_resume("broken.docx", data, max_bytes=LIMIT)
    assert result.status == ParseStatus.FAILED
    assert result.error_code == "corrupted_file"


def test_zip_that_is_not_a_docx():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    result = parse_resume("archive.docx", buf.getvalue(), max_bytes=LIMIT)
    assert result.status == ParseStatus.FAILED and result.error_code == "corrupted_file"


def test_empty_file():
    result = parse_resume("empty.pdf", b"", max_bytes=LIMIT)
    assert result.status == ParseStatus.EMPTY and result.error_code == "empty_file"


def test_empty_docx_and_blank_pdf():
    assert parse_resume("e.docx", make_docx([]), max_bytes=LIMIT).status == ParseStatus.EMPTY
    assert parse_resume("e.pdf", make_pdf([]), max_bytes=LIMIT).status == ParseStatus.EMPTY


def test_image_only_pdf_requires_ocr():
    result = parse_resume("scan.pdf", make_image_only_pdf(), max_bytes=LIMIT)
    assert result.status == ParseStatus.NEEDS_OCR
    assert result.error_code == "no_extractable_text"
    assert result.text is None


def test_password_protected_pdf():
    result = parse_resume("locked.pdf", make_pdf(password="s3cret"), max_bytes=LIMIT)
    assert result.status == ParseStatus.FAILED
    assert result.error_code == "password_protected"


def test_wrong_extension_is_processed_by_content_with_warning():
    result = parse_resume("resume.docx", make_pdf(), max_bytes=LIMIT)
    assert result.ok and result.detected_format == "pdf"
    assert any("does not match" in w for w in result.warnings)


def test_pdf_extension_with_non_pdf_content():
    result = parse_resume("resume.pdf", b"just some text pretending", max_bytes=LIMIT)
    assert result.status == ParseStatus.FAILED and result.error_code == "corrupted_file"


@pytest.mark.parametrize("name,data", [
    ("notes.txt", b"plain text resume"),
    ("photo.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 50),
    ("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 50),
])
def test_unsupported_formats(name, data):
    result = parse_resume(name, data, max_bytes=LIMIT)
    assert result.status == ParseStatus.FAILED and result.error_code == "unsupported_format"


def test_oversized_upload():
    result = parse_resume("big.pdf", make_pdf() + b"0" * LIMIT, max_bytes=LIMIT)
    assert result.status == ParseStatus.FAILED and result.error_code == "file_too_large"


def test_detect_format():
    assert detect_format(make_pdf()) == "pdf"
    assert detect_format(make_docx()) == "docx"
    assert detect_format(b"hello") is None


def test_filename_sanitisation():
    assert sanitize_filename("../../etc/passwd") == "passwd"
    assert sanitize_filename("C:\\Users\\x\\cv.pdf") == "cv.pdf"
    assert sanitize_filename("a\x00b.pdf") == "ab.pdf"
    assert sanitize_filename("") == "unnamed-file"
